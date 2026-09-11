from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
DOCX_DIR = ROOT / "docx-professional"


def dxa(value):
    return round(value.inches * 1440)


def all_text(doc):
    values = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def table_geometry_errors(table, index):
    errors = []
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None or tbl_w.get(qn("w:w")) != "9360":
        errors.append(f"table {index}: tblW != 9360")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None or tbl_ind.get(qn("w:w")) != "120":
        errors.append(f"table {index}: tblInd != 120")
    grid_widths = [int(c.get(qn("w:w"))) for c in table._tbl.tblGrid]
    if sum(grid_widths) != 9360:
        errors.append(f"table {index}: grid sum={sum(grid_widths)}")
    for row_index, row in enumerate(table.rows):
        widths = []
        for cell in row.cells:
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                errors.append(f"table {index} row {row_index}: missing tcW")
                continue
            widths.append(int(tc_w.get(qn("w:w"))))
        if widths and sum(widths) != 9360:
            errors.append(f"table {index} row {row_index}: cell sum={sum(widths)}")
    return errors


def verify(source, output):
    errors = []
    with zipfile.ZipFile(output) as package:
        bad = package.testzip()
        if bad:
            errors.append(f"broken ZIP member: {bad}")

    doc = Document(output)
    text = all_text(doc)
    source_text = source.read_text(encoding="utf-8")
    expected_title = source_text.splitlines()[0].removeprefix("# ").strip()
    if expected_title not in text:
        errors.append("title missing")
    for heading in re.findall(r"^##\s+(.+)$", source_text, re.MULTILINE):
        if heading not in text:
            errors.append(f"heading missing: {heading}")
    if "�" in text or "\ufeff" in text:
        errors.append("invalid/replacement Unicode character")
    if re.search(r"\b(gateway|portal|router)\b", text, re.IGNORECASE):
        errors.append("out-of-scope routing term found")
    if len(doc.paragraphs) < 35:
        errors.append(f"too few paragraphs: {len(doc.paragraphs)}")
    if len(doc.tables) < 2:
        errors.append(f"too few tables: {len(doc.tables)}")

    for section_index, section in enumerate(doc.sections):
        values = {
            "page_width": dxa(section.page_width),
            "page_height": dxa(section.page_height),
            "top_margin": dxa(section.top_margin),
            "bottom_margin": dxa(section.bottom_margin),
            "left_margin": dxa(section.left_margin),
            "right_margin": dxa(section.right_margin),
        }
        expected = {
            "page_width": 12240,
            "page_height": 15840,
            "top_margin": 1440,
            "bottom_margin": 1440,
            "left_margin": 1440,
            "right_margin": 1440,
        }
        for key, expected_value in expected.items():
            if abs(values[key] - expected_value) > 2:
                errors.append(f"section {section_index}: {key}={values[key]}")
        footer_xml = section.footer._element.xml
        if section_index > 0 and " PAGE " not in footer_xml:
            errors.append(f"section {section_index}: page field missing")

    for index, table in enumerate(doc.tables):
        errors.extend(table_geometry_errors(table, index))

    return {
        "file": output.name,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "errors": errors,
    }


def main():
    pairs = [
        (ROOT / "01-Tool-개발환경-가이드.md", DOCX_DIR / "01-Tool-개발환경-가이드.docx"),
        (ROOT / "02-Tool-설계-가이드.md", DOCX_DIR / "02-Tool-설계-가이드.docx"),
        (ROOT / "03-Tool-개발가이드-1차.md", DOCX_DIR / "03-Tool-개발가이드-1차.docx"),
    ]
    failed = False
    for source, output in pairs:
        result = verify(source, output)
        print(f"{result['file']}: paragraphs={result['paragraphs']}, tables={result['tables']}, sections={result['sections']}")
        for error in result["errors"]:
            failed = True
            print(f"  ERROR: {error}")
    if failed:
        raise SystemExit(1)
    print("DOCX structural QA passed")


if __name__ == "__main__":
    main()
