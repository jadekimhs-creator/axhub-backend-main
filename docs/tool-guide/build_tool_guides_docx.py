from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "docx-professional"

SOURCES = [
    (
        ROOT / "01-Tool-개발환경-가이드.md",
        "개발 환경 가이드",
        "Tool Pod 개발 환경과 내부망 반입 기준",
        ("개발 환경", "실행·검증", "내부망 반입"),
    ),
    (
        ROOT / "02-Tool-설계-가이드.md",
        "Tool 설계 가이드",
        "명명·구조·Schema·MCI·Glow HTTP 설계 기준",
        ("명명·패키지", "Input·Output Schema", "MCI·Glow HTTP"),
    ),
    (
        ROOT / "03-Tool-개발가이드-1차.md",
        "Tool 개발 가이드 1차본",
        "신규 Tool 구현·검증·완료 절차",
        ("Scaffold", "구현·연계", "테스트·완료"),
    ),
]

FONT_BODY = "맑은 고딕"
FONT_CODE = "Consolas"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
INK = "24364B"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
SOFT_BLUE = "F1F6FB"
TABLE_ALT = "F8FAFC"
BORDER = "C9D2DE"
WHITE = "FFFFFF"
GOLD = "B88928"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, name=FONT_BODY, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    if total != CONTENT_DXA:
        widths[-1] += CONTENT_DXA - total
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_bottom_border(paragraph, color=BORDER, size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_top_border(paragraph, color=BORDER, size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), size)
    top.set(qn("w:space"), "1")
    top.set(qn("w:color"), color)
    p_bdr.append(top)


def add_left_border(paragraph, color=BLUE, size="18"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Title": (30, NAVY, 0, 8),
        "Subtitle": (13.5, MUTED, 0, 18),
        "Heading 1": (16, NAVY, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = doc.styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_section_geometry(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    page_num = sect_pr.find(qn("w:pgNumType"))
    if page_num is None:
        page_num = OxmlElement("w:pgNumType")
        sect_pr.append(page_num)
    page_num.set(qn("w:start"), str(start))


def configure_cover_section(section):
    configure_section_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear_paragraph(section.header.paragraphs[0])
    clear_paragraph(section.footer.paragraphs[0])


def configure_content_section(section, short_title):
    configure_section_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    set_page_number_start(section, 1)

    header_p = section.header.paragraphs[0]
    clear_paragraph(header_p)
    header_p.paragraph_format.space_after = Pt(4)
    header_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = header_p.add_run(short_title)
    set_run_font(left, size=8.5, color=NAVY, bold=True)
    right = header_p.add_run("\tAX HUB · TOOL DEVELOPMENT STANDARD")
    set_run_font(right, size=8.2, color=MUTED)
    add_bottom_border(header_p, "D7DEE8", "4")

    footer_p = section.footer.paragraphs[0]
    clear_paragraph(footer_p)
    footer_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = footer_p.add_run("INTERNAL USE · TOOL POD DEVELOPMENT")
    set_run_font(left, size=8.2, color=MUTED)
    right = footer_p.add_run("\t")
    set_run_font(right, size=8.2, color=MUTED)
    add_page_number(footer_p)


def add_inline(paragraph, text, default_size=11, default_color=None):
    token_pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    cursor = 0
    for match in token_pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=default_color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=FONT_CODE, size=max(8.5, default_size - 1), color=DARK_BLUE)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F6")
            run._r.get_or_add_rPr().append(shd)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(label)
            set_run_font(run, size=default_size, color=BLUE)
            run.underline = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=default_size, color=default_color)


def add_cover(doc, title, subtitle, doc_type, highlights):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(88)
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("AX HUB  |  TOOL DEVELOPMENT STANDARD")
    set_run_font(r, size=9.5, color=GOLD, bold=True)

    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(10)
    add_inline(title_p, title, default_size=30, default_color=NAVY)

    subtitle_p = doc.add_paragraph(style="Subtitle")
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(subtitle_p, subtitle, default_size=13.5, default_color=DARK_BLUE)

    rule = doc.add_paragraph()
    rule.paragraph_format.left_indent = Inches(1.55)
    rule.paragraph_format.right_indent = Inches(1.55)
    rule.paragraph_format.space_after = Pt(24)
    add_bottom_border(rule, GOLD, "10")

    focus = doc.add_paragraph()
    focus.alignment = WD_ALIGN_PARAGRAPH.CENTER
    focus.paragraph_format.space_after = Pt(86)
    focus_run = focus.add_run("   ·   ".join(highlights))
    set_run_font(focus_run, size=10.5, color=BLUE, bold=True)

    rows = [
        ("문서 구분", doc_type),
        ("기준 버전", "1차본  |  2026-08-14"),
        ("적용 범위", "dap-was-lib 및 Tool Pod (CUS · SAL · PRO · SYS)"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(5)
        label_run = p.add_run(f"{label}  |  ")
        set_run_font(label_run, size=9.2, color=MUTED, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=9.5, color=INK)


def extract_section_titles(markdown_text):
    return [
        line.removeprefix("## ").strip()
        for line in markdown_text.splitlines()
        if line.startswith("## ")
    ]


def add_document_overview(doc, subtitle, markdown_text):
    heading = doc.add_heading("문서 안내", level=1)
    add_bottom_border(heading, LIGHT_BLUE, "8")

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(10)
    run = lead.add_run(subtitle)
    set_run_font(run, size=11.5, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), SOFT_BLUE)
    p_pr.append(shd)
    add_left_border(p, BLUE, "20")
    marker = p.add_run("READING GUIDE  ")
    set_run_font(marker, size=9, color=BLUE, bold=True)
    body = p.add_run("이 문서는 Tool Pod 개발자가 설계·구현·검증 과정에서 바로 참고할 수 있도록 현재 소스 기준으로 정리했습니다.")
    set_run_font(body, size=10.2, color=INK)

    section_heading = doc.add_heading("주요 구성", level=2)
    section_heading.paragraph_format.space_after = Pt(6)
    for title in extract_section_titles(markdown_text):
        item = doc.add_paragraph(style="List Bullet")
        add_inline(item, title, default_size=10.2, default_color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(4)
    add_bottom_border(rule, "D7DEE8", "4")


def is_table_separator(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        if not is_table_separator(lines[idx]):
            rows.append([c.strip() for c in lines[idx].strip().strip("|").split("|")])
        idx += 1
    return rows, idx


def table_widths(rows):
    col_count = max(len(row) for row in rows)
    weights = []
    for col in range(col_count):
        max_len = max(len(row[col]) if col < len(row) else 0 for row in rows)
        weights.append(max(8, min(max_len, 48)))
    total = sum(weights)
    widths = [max(900, round(CONTENT_DXA * weight / total)) for weight in weights]
    scale = CONTENT_DXA / sum(widths)
    widths = [round(width * scale) for width in widths]
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_markdown_table(doc, rows):
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for row_idx, values in enumerate(rows):
        for col_idx in range(col_count):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            value = values[col_idx] if col_idx < len(values) else ""
            add_inline(p, value, default_size=8.8 if col_count >= 4 else 9.2)
            if row_idx == 0:
                set_cell_shading(cell, NAVY)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, TABLE_ALT)
            if row_idx > 0:
                column_values = [row[col_idx] if col_idx < len(row) else "" for row in rows[1:]]
                if column_values and max(len(value) for value in column_values) <= 12:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repeat_table_header(table.rows[0])
    set_table_geometry(table, table_widths(rows))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc, code_lines):
    if not code_lines:
        return
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = len(code_lines) <= 12
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F5F7FA")
    p_pr.append(shd)
    add_left_border(p, BLUE, "16")
    for index, line in enumerate(code_lines):
        run = p.add_run(line)
        set_run_font(run, name=FONT_CODE, size=8.2, color=NAVY)
        if index < len(code_lines) - 1:
            run.add_break()


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), SOFT_BLUE)
    p_pr.append(shd)
    add_left_border(p)
    marker = p.add_run("NOTE  ")
    set_run_font(marker, size=9.5, color=BLUE, bold=True)
    add_inline(p, text, default_size=10)


def add_markdown_body(doc, markdown_text):
    lines = markdown_text.splitlines()
    first_h1_seen = False
    in_code = False
    code_lines = []
    idx = 0
    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue
        if not stripped:
            idx += 1
            continue
        if stripped.startswith("# ") and not first_h1_seen:
            first_h1_seen = True
            idx += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            idx += 1
            continue
        if stripped.startswith("## "):
            heading = doc.add_heading(stripped[3:], level=1)
            add_bottom_border(heading, LIGHT_BLUE, "8")
            idx += 1
            continue
        if stripped.startswith("# "):
            heading = doc.add_heading(stripped[2:], level=1)
            add_bottom_border(heading, LIGHT_BLUE, "8")
            idx += 1
            continue
        if stripped.startswith("|") and idx + 1 < len(lines) and is_table_separator(lines[idx + 1]):
            rows, idx = parse_table(lines, idx)
            add_markdown_table(doc, rows)
            continue
        if stripped.startswith("> "):
            add_callout(doc, stripped[2:])
            idx += 1
            continue
        match_ordered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if match_ordered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, match_ordered.group(1))
            idx += 1
            continue
        if stripped.startswith("- [ ] "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, "☐ " + stripped[6:])
            idx += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            idx += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped)
        idx += 1
    if code_lines:
        add_code_block(doc, code_lines)


def set_document_properties(doc, title):
    props = doc.core_properties
    props.title = title
    props.subject = "AX HUB Tool 개발 표준"
    props.author = "AX HUB MCP & TOOL"
    props.keywords = "AX HUB, Tool Pod, MCP, Glow, MCI, HTTP"
    props.comments = "Tool Pod 개발자용 문서"


def build(source_path, short_title, subtitle, highlights):
    markdown = source_path.read_text(encoding="utf-8")
    first_line = markdown.splitlines()[0]
    full_title = first_line.removeprefix("# ").strip()

    doc = Document()
    configure_styles(doc)
    set_document_properties(doc, full_title)
    configure_cover_section(doc.sections[0])
    add_cover(doc, full_title, subtitle, short_title, highlights)

    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_content_section(content_section, short_title)
    add_document_overview(doc, subtitle, markdown)
    add_markdown_body(doc, markdown)

    output = OUTPUT_DIR / f"{source_path.stem}.docx"
    doc.save(output)
    return output


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for source, short_title, subtitle, highlights in SOURCES:
        output = build(source, short_title, subtitle, highlights)
        print(output)


if __name__ == "__main__":
    main()
